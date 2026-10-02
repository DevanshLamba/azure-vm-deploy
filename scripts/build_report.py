"""Build docs/report.docx and docs/report.pdf from docs/report.md.

Handles the Markdown used in the report: headings, paragraphs, bullet/numbered lists (one nested
level), tables, code blocks, block quotes ("PENDING" ones are highlighted), images with a
"*Figure N: ...*" caption line, **bold**, *italic* and `code`. The Mermaid diagram is replaced
by docs/architecture.png (the same Mermaid source rendered top-to-bottom, which fits a portrait
page better). Microsoft Word (via COM) then fills in the table of contents and page
numbers and exports the PDF.

Usage: python scripts/build_report.py
"""
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

DOCS = Path(__file__).resolve().parent.parent / "docs"
MD = DOCS / "report.md"
DOCX = DOCS / "report.docx"
PDF = DOCS / "report.pdf"

INK = RGBColor(0x2F, 0x2A, 0x33)
ACCENT = RGBColor(0x5B, 0x4B, 0xA8)
MUTED = RGBColor(0x5C, 0x55, 0x62)
CONTENT_WIDTH = 6.3  # inches (A4 with 1" margins, roughly)

TITLE = "Deployment to a Publicly Hosted Linux VM on Azure"
SUBTITLE = "CloudTasks: a task tracker with a live VM status panel"
STUDENT = "Devansh Lamba"
COLLEGE = "Bharati Vidyapeeth College of Engineering, Pune"
SUBJECT = "Cloud Computing"
PROJECT = "PBL project #16"


# ------------------------------------------------------------------ low-level helpers
def shade(element, hex_fill):
    pr = element.get_or_add_pPr() if hasattr(element, "get_or_add_pPr") else element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def add_field(paragraph, instruction, placeholder=""):
    """Insert a Word field (TOC, PAGE, NUMPAGES). Word computes its value when fields are updated."""
    def fld(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el
    run = paragraph.add_run()
    run._r.append(fld("begin"))
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)
    run._r.append(fld("separate"))
    result = paragraph.add_run(placeholder)
    end = paragraph.add_run()
    end._r.append(fld("end"))
    return result, run, end


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*\*|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text, size=None, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt((size or 11) - 1.5)
            run.font.color.rgb = ACCENT
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            run = paragraph.add_run(m.group(1))
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)
        if color is not None and run.font.color.rgb is None:
            run.font.color.rgb = color
    return paragraph


# ------------------------------------------------------------------ document setup
def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.font.color.rgb = INK
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Title", 26)):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = INK if name != "Heading 2" else ACCENT
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(18)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(12)

    sec = doc.sections[0]
    sec.page_height, sec.page_width = Inches(11.69), Inches(8.27)  # A4
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(1.0))
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    sec.different_first_page_header_footer = True  # no page number on the title page
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text, field in (("Page ", "PAGE"), (" of ", "NUMPAGES")):
        r = foot.add_run(text)
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED
        for fr in add_field(foot, field, "1"):
            fr.font.size = Pt(9)
            fr.font.color.rgb = MUTED
    hdr = sec.header.paragraphs[0]
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hdr.add_run(f"{SUBJECT} · {PROJECT} · {STUDENT}")
    r.font.size = Pt(8.5)
    r.font.color.rgb = MUTED

    props = doc.core_properties
    props.title = TITLE
    props.author = STUDENT
    props.subject = f"{SUBJECT}, {PROJECT}"


def title_page(doc):
    def centered(text, size, bold=False, color=INK, space_after=6, italic=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = color
        return p

    centered(COLLEGE, 15, bold=True, space_after=2)
    centered(SUBJECT, 11, color=MUTED, space_after=60)
    centered(f"Project Based Learning (PBL), {PROJECT.split(' ', 1)[1].capitalize()}", 12, color=ACCENT, space_after=10)
    centered(TITLE, 26, bold=True, space_after=8)
    centered(SUBTITLE, 13, italic=True, color=MUTED, space_after=50)

    centered("Submitted by", 11, bold=True, space_after=6)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ("#", "Name", "Roll number")):
        cell.text = ""
        r = cell.paragraphs[0].add_run(text)
        r.bold = True
        shade(cell._tc, "E7DFFB")
    rows = [(STUDENT, ""), ("", ""), ("", ""), ("", "")]
    for i, (name, roll) in enumerate(rows, 1):
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = str(i), name or " ", roll or " "
    for row in table.rows:
        row.height = Inches(0.36)
        for cell, width in zip(row.cells, (0.4, 3.0, 2.0)):
            cell.width = Inches(width)

    doc.add_paragraph().paragraph_format.space_after = Pt(40)
    centered("Subject: " + SUBJECT, 11, space_after=2)
    centered("Guide: ______________________________", 11, color=MUTED, space_after=2)
    centered("October 2026", 11, color=MUTED)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    h = doc.add_paragraph()
    r = h.add_run("Contents")
    r.bold = True
    r.font.size = Pt(18)
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-2" \\h \\z \\u', "Table of contents: open in Word and press F9 to update.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------ blocks
def add_image(doc, rel_path, caption):
    path = (DOCS / rel_path).resolve()
    w, h = Image.open(path).size
    width = CONTENT_WIDTH if h / w < 1.2 else 2.9  # portrait (mobile) shots narrower
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_caption(c, caption)


def add_caption(paragraph, text):
    """'*Figure 1: ... `code` ...*' -> italic grey caption that still renders inline code."""
    add_inline(paragraph, text.strip().strip("*"), size=9.5, color=MUTED)
    for r in paragraph.runs:
        r.italic = True


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], [r for r in cells[2:]]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = ""
        add_inline(cell.paragraphs[0], text, size=9.5)
        for r in cell.paragraphs[0].runs:
            r.bold = True
        shade(cell._tc, "E7DFFB")
    for row in body:
        rc = t.add_row().cells
        for cell, text in zip(rc, row):
            cell.text = ""
            add_inline(cell.paragraphs[0], text, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(doc, lines):
    for i, line in enumerate(lines or [""]):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(4) if i == 0 else Pt(0)
        pf.line_spacing = 1.0
        pf.keep_with_next = i < len(lines) - 1
        shade(p._p, "F5EFE9")
        r = p.add_run(line.rstrip() or " ")
        r.font.name = "Consolas"
        r.font.size = Pt(6.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_quote(doc, text):
    pending = text.startswith("**PENDING")
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    shade(p._p, "FFE9A8" if pending else "F1ECF8")
    add_inline(p, text, size=10.5)
    return p


def build():
    lines = MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    setup(doc)
    title_page(doc)

    i = 0
    para = []

    def flush():
        if para:
            add_inline(doc.add_paragraph(), " ".join(para))
            para.clear()

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if raw.startswith("# "):  # document title: already on the title page
            i += 1
            continue
        if line.startswith("```"):
            flush()
            lang = line[3:].strip()
            block = []
            i += 1
            while not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            if lang == "mermaid":
                add_image(doc, "architecture.png", "*Figure: Architecture of the deployment (rendered from the Mermaid diagram).*")
            else:
                add_code(doc, block)
            continue
        if line.startswith("## "):
            flush()
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            flush()
            doc.add_heading(line[4:], level=2)
        elif line == "---" or not line:
            flush()
        elif line.startswith("!["):
            flush()
            m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", line)
            caption = None
            if i + 1 < len(lines) and lines[i + 1].strip().startswith("*Figure"):
                caption = lines[i + 1].strip()
                i += 1
            add_image(doc, m.group(1), caption)
        elif line.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        elif line.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            if quote[0].startswith("**PENDING") and len(quote) > 1:  # multi-line: keep the list shape
                add_quote(doc, quote[0])
                for q in quote[1:]:
                    p = add_quote(doc, q)
                    p.paragraph_format.left_indent = Inches(0.45)
            else:
                add_quote(doc, " ".join(quote))
            if i < len(lines) and lines[i].strip().startswith("*Figure"):  # caption of a pending figure
                add_caption(doc.add_paragraph(), lines[i])
                i += 1
            continue
        elif re.match(r"^(\s*)([-*]|\d+\.)\s+", raw):
            flush()
            m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", raw)
            nested = len(m.group(1)) >= 2
            numbered = m.group(2)[0].isdigit()
            style = ("List Number" if numbered else "List Bullet") + (" 2" if nested else "")
            add_inline(doc.add_paragraph(style=style), m.group(3))
        elif line.startswith("*Listing") or line.startswith("*Figure"):
            flush()
            c = doc.add_paragraph()
            c.paragraph_format.keep_with_next = True
            add_caption(c, line)
        else:
            para.append(line)
        i += 1
    flush()
    doc.save(DOCX)
    print("wrote", DOCX.relative_to(DOCS.parent))


def word_finish():
    """Let Word fill in the table of contents and page numbers, save, and export the PDF."""
    ps = f"""
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false; $w.DisplayAlerts = 0
try {{
  $d = $w.Documents.Open('{DOCX}')
  $d.TablesOfContents.Item(1).Update()
  $null = $d.Fields.Update()
  $d.TablesOfContents.Item(1).UpdatePageNumbers()
  $d.Save()
  # 17 = PDF; CreateBookmarks 1 = from headings (PDF outline)
  $d.ExportAsFixedFormat('{PDF}', 17, $false, 0, 0, 1, 1, 0, $true, $true, 1, $true, $true, $false)
  $d.Close()
}} finally {{ $w.Quit() }}
"""
    subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps], check=True)
    print("updated TOC + page numbers in", DOCX.relative_to(DOCS.parent), "and wrote", PDF.relative_to(DOCS.parent))


if __name__ == "__main__":
    build()
    word_finish()
